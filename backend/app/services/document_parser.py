"""Document parsing service supporting PDF, DOCX, and scanned images.
Implements pdfplumber, python-docx, and Gemini multimodal OCR fallback.
"""
import io
import os
import logging
from typing import Tuple, Dict, Any, List

logger = logging.getLogger("lexguard.parser")

class DocumentParserService:
    """Parses legal documents across text-based and scanned image formats."""

    @staticmethod
    def validate_file(filename: str, file_size: int, max_size_mb: int = 15) -> Tuple[bool, str]:
        """Validate filename extension and file size limit."""
        if not filename or "." not in filename:
            return False, "File must have a valid extension."

        ext = os.path.splitext(filename)[1].lower()
        allowed = [".pdf", ".docx", ".png", ".jpg", ".jpeg"]
        if ext not in allowed:
            return False, f"Unsupported file format '{ext}'. Allowed: {', '.join(allowed)}"

        max_bytes = max_size_mb * 1024 * 1024
        if file_size > max_bytes:
            return False, f"File exceeds maximum size limit of {max_size_mb}MB."

        return True, ""

    @classmethod
    async def extract_content(cls, file_bytes: bytes, filename: str, gemini_service: Any = None) -> Dict[str, Any]:
        """Extract text, page mapping, and metadata from document bytes."""
        ext = os.path.splitext(filename)[1].lower()
        extracted_text = ""
        page_texts: List[str] = []
        is_scanned = False
        method_used = "unknown"

        if ext == ".pdf":
            extracted_text, page_texts, is_scanned, method_used = await cls._parse_pdf(file_bytes, gemini_service)
        elif ext == ".docx":
            extracted_text, page_texts, method_used = cls._parse_docx(file_bytes)
        elif ext in [".png", ".jpg", ".jpeg"]:
            extracted_text, page_texts, method_used = await cls._parse_image(file_bytes, ext, gemini_service)
            is_scanned = True
        else:
            raise ValueError(f"Unsupported file format: {ext}")

        return {
            "filename": filename,
            "file_type": ext.replace(".", "").upper(),
            "raw_text": extracted_text.strip(),
            "pages": page_texts,
            "page_count": len(page_texts),
            "is_scanned": is_scanned,
            "extraction_method": method_used,
            "character_count": len(extracted_text)
        }

    @classmethod
    async def _parse_pdf(cls, file_bytes: bytes, gemini_service: Any) -> Tuple[str, List[str], bool, str]:
        """Parse PDF using pdfplumber, falling back to Gemini multimodal OCR if scanned."""
        page_texts = []
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for idx, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    page_texts.append(text)
        except Exception as e:
            logger.warning("pdfplumber extraction failed: %s. Attempting raw text decode.", str(e))
            try:
                decoded = file_bytes.decode("utf-8", errors="ignore")
                if len(decoded.strip()) > 20:
                    page_texts = [decoded]
            except Exception:
                pass

        combined_text = "\n\n".join(page_texts).strip()

        # If combined text is empty or suspiciously short (<60 chars across multiple pages), it's likely scanned
        is_scanned = len(combined_text) < 60
        if is_scanned and gemini_service:
            logger.info("PDF has sparse/empty text layer. Triggering Gemini Multimodal OCR fallback...")
            try:
                ocr_prompt = (
                    "Transcribe the full legal document from this scanned PDF exactly as written. "
                    "Preserve clause titles, numbers, and paragraph breaks. Do not summarize; provide verbatim transcription."
                )
                transcribed = await gemini_service.analyze_multimodal(file_bytes, "application/pdf", ocr_prompt)
                if transcribed and len(transcribed.strip()) > 50:
                    return transcribed.strip(), [transcribed.strip()], True, "gemini_multimodal_ocr"
            except Exception as e:
                logger.error("Gemini OCR fallback failed: %s", str(e))

        method = "pdfplumber" if not is_scanned else "pdfplumber_low_fidelity"
        return combined_text, page_texts if page_texts else [combined_text], is_scanned, method

    @classmethod
    def _parse_docx(cls, file_bytes: bytes) -> Tuple[str, List[str], str]:
        """Extract text from DOCX file using python-docx, falling back to text decoding."""
        paragraphs = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    paragraphs.append(text)

            # Also extract table text
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            method = "python-docx"
        except Exception as e:
            logger.warning("python-docx extraction failed: %s. Attempting raw text decode fallback.", str(e))
            decoded = file_bytes.decode("utf-8", errors="ignore")
            paragraphs = [decoded] if decoded.strip() else []
            method = "text_decode_fallback"

        combined_text = "\n\n".join(paragraphs)
        return combined_text, [combined_text], method

    @classmethod
    async def _parse_image(cls, file_bytes: bytes, ext: str, gemini_service: Any) -> Tuple[str, List[str], str]:
        """Extract legal text from image using Gemini multimodal vision."""
        mime_type = "image/png" if ext == ".png" else "image/jpeg"
        if gemini_service:
            ocr_prompt = (
                "Transcribe all contractual and legal clauses from this image clearly and completely. "
                "Retain section headers, numbers, and full clause texts without truncation."
            )
            transcribed = await gemini_service.analyze_multimodal(file_bytes, mime_type, ocr_prompt)
            if transcribed:
                return transcribed.strip(), [transcribed.strip()], "gemini_vision_ocr"

        # Fallback if no service is configured
        return "Image document processed. (Gemini Vision OCR API key required for full image transcription).", ["Image document."], "vision_fallback"

document_parser = DocumentParserService()
